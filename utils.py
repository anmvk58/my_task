from openpyxl import load_workbook
import glob, os
import xlsxwriter

result_temp = []

def get_data(file_path):
    wb = load_workbook(filename=file_path)
    # get name of sheet
    sheet_name_table = wb.sheetnames[-1]

    result = []
    # get data of sheet
    ws = wb[sheet_name_table]
    for x in range(ws.max_row):
        if x == 0:
            continue
        else:
            o_cell = ws.cell(row=x + 1, column=1)
            result.append(o_cell.value)

    return result

def get_list_file(folder_path):
    os.chdir(folder_path)
    list_file = glob.glob("*.dtsx")
    return list_file

def write_result_to_xlsx(file_path, data):
    # Create new file to write
    workbook = xlsxwriter.Workbook((''.join([file_path,"\\" ,'Result.xlsx'])))
    worksheet = workbook.add_worksheet()

    # Define format to apply somewhere
    cell_format = workbook.add_format({'bold': True})
    format_red = workbook.add_format({'bg_color': '#FFC7CE', 'font_color': '#9C0006'})
    format_green = workbook.add_format({'bg_color': '#C6EFCE', 'font_color': '#006100'})

    # Set width for column
    worksheet.set_column(0, 0, 165)
    worksheet.set_column(1, 1, 40)

    # Write header of column
    worksheet.write(0, 0, 'File_Path', cell_format)
    worksheet.write(0, 1, 'Table Name', cell_format)

    row = 1
    for path, key in (data):
        worksheet.write(row, 0, path)
        worksheet.write(row, 1, key)
        row += 1

    # Add highlight rule for path
    range = 'A2:A' + str(row - 1)
    worksheet.conditional_format(range, {'type': 'duplicate', 'format': format_red})

    workbook.close()


if __name__ == '__main__':
    # #test get table name to search
    # result = get_data('C:\\Users\\anmv1\\Desktop\\New folder\\LIST_TABLE_TMP.xlsx')
    # for cell in result:
    #     print(cell)

    #test get file in folder
    result = get_list_file("D:\\Source\\SSIS\\SSIS_srv20\\ETL_G01_TRACKING\\ETL_G01_TRACKING\\VPB_ETL_TRACKING")
